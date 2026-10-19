"""网页看板：实时状态 + 备忘录卡片流 + 一键导出。

设计要点
--------
* **Flask 是懒加载的**：所有 ``from flask import ...`` 都写在 ``create_app()`` 内部。
  没装 Flask 时 ``import webui`` 依然成功，只是 ``create_app`` 抛 ImportError ——
  ``main.py`` 捕获后给一句安装提示，其余功能照常。自检脚本也靠这一点优雅跳过。
* **页面不轮询**：``/api/stream`` 是一条 SSE（``text/event-stream``）长连接，
  状态变化和"新卡片落库"由 ``main.Broadcaster`` 主动推过来，页面收到即渲染。
* **写入只走 main 里的函数**：转写、整理、落库全在 ``main.MemoPipeline`` 里；删除走
  ``main.remove_memo``（顺带把 ``my_memos.md`` 重刷对齐）。Web 层自己不写 SQL，
  只是那套状态机的一面显示器 + 一个删除按钮。

路由
----
``GET  /``                 卡片看板页面
``GET  /api/state``        当前状态快照（页面首次加载时用）
``GET  /api/memos``        最近的卡片（按时间倒序）
``GET  /api/stream``       SSE 实时推送（state / memo）
``POST /api/text``         手动塞一段文字进整理链路（--no-mic 时很好用）
``POST /api/stop``         停止监听（等同于语音说"退出录制"）
``POST /api/control/finish``  手动结束并整理当前录音（说不出指令词时的双保险）
``POST /api/control/cancel``  手动放弃当前录音（不转写、不入库）
``POST /api/memos/update/<id>``        💾 保存修改：contenteditable 回写（数据库 + my_memos.md 同步）
``POST /api/memos/re-summarize/<id>``  🔄 智能润色/推演：交给 DeepSeek 做临床推理后写回（写完推 announce 事件）
``POST /api/memos/merge``              🔗 多选合并：把勾选的几段记录熔成一张新病历卡片（合并后推 announce 事件）
``POST|DELETE /api/delete/<id>``  删除一张卡片（数据库 + my_memos.md 同步）

朗读不在后端做：整理 / 润色 / 合并完成后，后端只通过 SSE 推一条 ``announce``
事件（带整张卡片），由页面用浏览器原生的 Web Speech API 出声。
``POST /api/memos/<id>/speak`` 与 ``/api/tts*`` 已经移除，避免再引入 edge-tts。
``GET  /api/export``       导出全部记录：``?fmt=md`` 或 ``?fmt=docx``，直接下载

上面两个 ``/api/memos/...`` 是现行路径；早先的 ``/api/memo/<id>/update`` 与
``/api/memo/<id>/regenerate`` 保留为**别名**（同一个视图函数挂两个 url rule），
老书签、老脚本、以及浏览器里缓存的旧页面都还能用。
"""

from __future__ import annotations

import io
import sqlite3
import time
from datetime import datetime
from typing import Any, Optional

DEFAULT_EXPORT_LIMIT = 500

#: 合并至少要两条 —— 一条的话"合并"就等于"重新润色"，学生该点的是那张卡上的 🔄。
#: 少给一条就报错，而不是默默按一条处理：不然他会以为合并过了，其实只是原地润了一遍。
MERGE_MIN_SOURCES = 2

#: 一次最多合并几条。挡的不是恶意请求，是"全选 50 条"这种手滑：
#: 输入会撑爆上下文、账单也跟着涨，而 50 段熔出来的病历本来也没法看。
MERGE_MAX_SOURCES = 20


def create_app(cfg: Any, conn: sqlite3.Connection, state: Any,
               pipeline: Any = None, listener: Any = None) -> Any:
    """构建并返回 Flask app。

    ``cfg``      : ``main.Config``
    ``conn``     : 已打开的 sqlite3 连接（``check_same_thread=False``）
    ``state``    : ``main.AppState``，提供 ``snapshot()``、``broadcaster``
                   与 ``publish_announce()``（朗读事件由它推给页面）
    ``pipeline`` : ``main.MemoPipeline``，供 ``/api/text`` 手动写入
    ``listener`` : ``main.VoiceListener``，供 ``/api/stop`` 停录音
    """
    try:
        from flask import Flask, Response, jsonify, render_template, request, send_file
    except ImportError as exc:  # pragma: no cover - 取决于是否装了 Flask
        raise ImportError(
            "网页看板需要 Flask，请先安装：pip install flask"
        ) from exc

    import main as core          # 复用 main.py 里的导出/查询函数，避免逻辑重复

    app = Flask(__name__, template_folder="templates", static_folder="static")
    app.config["JSON_AS_ASCII"] = False
    app.config["STATE"] = state
    app.config["LISTENER"] = listener
    # 重新润色要用它身上的 DeepSeek 客户端（可能为 None —— 没 Key 或 --no-llm）
    app.config["PIPELINE"] = pipeline

    # ------------------------------------------------------------------ #
    # 页面
    # ------------------------------------------------------------------ #
    @app.get("/")
    def index() -> Any:
        snapshot = state.snapshot()
        return render_template(
            "index.html",
            snapshot=snapshot,
            memos=[core.memo_to_dict(m) for m in core.fetch_memos(conn, DEFAULT_EXPORT_LIMIT)],
            segment_end=" / ".join(cfg.segment_end_keywords[:3]),
            exit_words=" / ".join(cfg.exit_keywords[:2]),
            cancel_words=" / ".join(cfg.cancel_keywords[:2]),
            silence_s=f"{cfg.silence_ms / 1000:.1f}",
        )

    # ------------------------------------------------------------------ #
    # JSON 接口
    # ------------------------------------------------------------------ #
    @app.get("/api/state")
    def api_state() -> Any:
        return jsonify(state.snapshot())

    @app.get("/api/memos")
    def api_memos() -> Any:
        try:
            limit = int(request.args.get("limit", DEFAULT_EXPORT_LIMIT))
        except (TypeError, ValueError):
            limit = DEFAULT_EXPORT_LIMIT
        limit = max(1, min(limit, 5000))
        memos = [core.memo_to_dict(m) for m in core.fetch_memos(conn, limit)]
        return jsonify({"count": len(memos), "memos": memos})

    @app.get("/api/stream")
    def api_stream() -> Any:
        """SSE：状态变化与新卡片实时推给页面。"""

        def events() -> Any:
            box = state.broadcaster.subscribe()
            try:
                # 刚连上先补一份当前状态，免得页面在第一次变化前空着
                yield _sse(state.published_state())
                while True:
                    try:
                        event = box.get(timeout=15)
                    except Exception:            # queue.Empty
                        yield ": keep-alive\n\n"  # 心跳，防止代理掐断长连接
                        continue
                    yield _sse(event)
            finally:
                state.broadcaster.unsubscribe(box)

        return Response(
            events(),
            mimetype="text/event-stream",
            headers={
                "Cache-Control": "no-cache, no-transform",
                "X-Accel-Buffering": "no",       # 关掉 nginx 缓冲（若前面挂了代理）
                "Connection": "keep-alive",
            },
        )

    @app.post("/api/text")
    def api_text() -> Any:
        """手动喂一段文字，走完整的整理 + 落库 + 推送链路。"""
        payload = request.get_json(silent=True) or {}
        text = (payload.get("text") or request.form.get("text") or "").strip()
        if not text:
            return jsonify({"ok": False, "error": "text 不能为空"}), 400
        if pipeline is None:
            return jsonify({"ok": False, "error": "当前进程没有可用的整理管线"}), 503
        try:
            memo_id = pipeline.handle_text(core.clean_transcript(text, cfg))
        except Exception as exc:                                  # 不让界面 500
            return jsonify({"ok": False, "error": f"{type(exc).__name__}: {exc}"}), 500
        if memo_id is None:
            return jsonify({"ok": False, "error": "内容只剩指令词，已忽略"}), 200
        return jsonify({"ok": True, "id": memo_id})

    @app.route("/api/delete/<int:memo_id>", methods=["POST", "DELETE"])
    def api_delete(memo_id: int) -> Any:
        """删除一张卡片：先删数据库，再按数据库现状重刷 my_memos.md，最后广播。

        删完必须广播 —— 否则同一个后端的其它标签页/手机页面还挂着那张已经不存在
        的卡片，点进去导出的文档却已经没有它了。
        """
        try:
            removed = core.remove_memo(conn, memo_id, cfg.markdown_path)
        except Exception as exc:                                  # 不让界面 500
            return jsonify({"ok": False, "error": f"{type(exc).__name__}: {exc}"}), 500
        if removed is None:
            return jsonify({"ok": False, "error": f"没有 id={memo_id} 的卡片"}), 404
        state.publish_delete(memo_id)
        remaining = conn.execute("SELECT COUNT(*) FROM memos").fetchone()[0]
        return jsonify({
            "ok": True,
            "id": memo_id,
            "title": removed.get("title") or "",
            "remaining": int(remaining),
            "markdown_path": cfg.markdown_path or "",
            "markdown_synced": bool(cfg.markdown_path),
        })

    def _read_payload() -> tuple:
        """解析两个写接口共用的请求体，返回 ``(fields, touched)``。

        ``fields``  : 页面上每个可编辑字段的当前文本（字典）
        ``touched`` : 学生**实际打过字**的字段名（列表），只有这些受"润色不许改回去"
                      的保护；``None`` 表示请求里压根没这一项。

        ``touched`` 缺失（浏览器里缓存的旧页面）时返回 ``None``，由调用方按"全都当改过"
        处理：旧页面报上来的 fields 分不出"改过"和"只是有值"，那就一律保护 ——
        少润色几段只是效果打折，拿模型结果盖掉学生核对过的描述却是真丢东西。
        """
        payload = request.get_json(silent=True) or {}
        raw = payload.get("fields") if isinstance(payload.get("fields"), dict) else payload
        fields = raw if isinstance(raw, dict) else {}
        touched = payload.get("touched")
        if not isinstance(touched, (list, tuple)):
            return fields, None
        return fields, [str(name) for name in touched]

    @app.post("/api/memos/update/<int:memo_id>")
    @app.post("/api/memo/<int:memo_id>/update")      # 旧路径，保留兼容
    def api_memo_update(memo_id: int) -> Any:
        """💾 保存修改：保存网页上手工修改的字段（``contenteditable`` 回写）。

        字段名走 ``core.EDITABLE_FIELDS`` 白名单：它既是"哪些字段允许改"的声明，
        也是拼 UPDATE 语句时唯一的列名来源 —— 白名单之外的一律丢弃，绝不拼进 SQL。
        """
        fields, _touched = _read_payload()
        if not fields:
            return jsonify({"ok": False, "error": "请求体得是字段字典"}), 400
        unknown = [name for name in fields if name not in core.EDITABLE_FIELDS]
        if unknown and not (set(fields) & set(core.EDITABLE_FIELDS)):
            return jsonify({
                "ok": False,
                "error": f"没有可改的字段；允许：{', '.join(core.EDITABLE_FIELDS)}",
            }), 400
        try:
            memo = core.update_memo(conn, memo_id, fields)
        except Exception as exc:                                  # 不让界面 500
            return jsonify({"ok": False, "error": f"{type(exc).__name__}: {exc}"}), 500
        if memo is None:
            return jsonify({"ok": False, "error": f"没有 id={memo_id} 的卡片"}), 404
        synced = _sync_markdown_and_broadcast(memo)
        saved = [name for name in fields if name in core.EDITABLE_FIELDS]
        return jsonify({
            "ok": True,
            "id": memo_id,
            "memo": core.memo_to_dict(memo),
            "saved": saved,
            "ignored": unknown,
            "markdown_path": cfg.markdown_path or "",
            "markdown_synced": synced,
        })

    @app.post("/api/memos/re-summarize/<int:memo_id>")
    @app.post("/api/memo/<int:memo_id>/regenerate")  # 旧路径，保留兼容
    def api_memo_regenerate(memo_id: int) -> Any:
        """🔄 智能润色 / 诊断推演：把（手改过的）整张卡片交给 DeepSeek 做临床推理。

        和"保存修改"是两回事 —— 保存是学生自己动手，这里是让模型像带教查房那样，
        基于他问到的信息推出**诊断思考**与**诊疗计划**（见 ``core.CLINICAL_PROMPT``）。

        两道保护：
        * 本次会话里学生亲手动过的字段，模型返回后被**强制还原**（``restore_edited_fields``），
          "忠于学生的修改"因此是代码保证的，不是提示词里的一句请求；
        * 写库发生在模型返回**之后** —— 调用失败时页面上的改动仍是"未保存"状态，
          学生不会误以为已经存进去了。
        """
        memo = core.fetch_memo(conn, memo_id)
        if memo is None:
            return jsonify({"ok": False, "error": f"没有 id={memo_id} 的卡片"}), 404
        # 页面上的未保存改动一起带过来：推演要基于"学生改过的那一版"而不是库里的旧版。
        fields, touched = _read_payload()
        if fields:
            memo = dict(memo)
            memo.update({name: str(value or "") for name, value in fields.items()
                         if name in core.EDITABLE_FIELDS})
        target = app.config.get("PIPELINE")
        client = getattr(target, "client", None)
        if client is None:
            return jsonify({
                "ok": False,
                "error": "当前没有可用的 DeepSeek 客户端（没配 API Key，或用了 --no-llm），"
                         "智能润色与诊断推演需要它",
            }), 503
        # touched 缺失（旧页面）时保守行事：页面上报的字段全都当成"他改过的"
        protect = list(fields) if touched is None else touched
        started = time.time()
        try:
            fresh = core.regenerate_memo(client, cfg.llm_model, memo,
                                         memo.get("record_time") or "")
            restored = core.restore_edited_fields(fresh, memo, protect)
            updated = core.update_memo(conn, memo_id, fresh,
                                       allowed=core.REGENERATE_FIELDS)
        except Exception as exc:                                  # 网络/解析失败都走这里
            return jsonify({"ok": False, "error": f"{type(exc).__name__}: {exc}"}), 502
        if updated is None:
            return jsonify({"ok": False, "error": f"没有 id={memo_id} 的卡片"}), 404
        synced = _sync_markdown_and_broadcast(updated)
        # 润色完让页面念一遍新的诊断与计划（浏览器 Web Speech，失败不影响本接口）
        state.publish_announce(updated, "regenerate")
        return jsonify({
            "ok": True,
            "id": memo_id,
            "memo": core.memo_to_dict(updated),
            # 如实告诉页面：哪几段是学生自己写的、被原样保住了（界面上会标出来）
            "restored": restored,
            "elapsed_ms": int((time.time() - started) * 1000),
            "markdown_path": cfg.markdown_path or "",
            "markdown_synced": synced,
        })

    def _parse_memo_ids(payload: dict) -> tuple:
        """把页面报上来的 ``memo_ids`` 规整成"去重后的 int 列表"，返回 ``(ids, error)``。

        去重是有意义的：页面上两张卡勾不出重复 id，手写 curl 会，
        而重复 id 意味着同一段记录在输入里出现两次 —— 模型据此可能写成
        "两次记录均提示……"，凭空多出一个并不存在的病程节点。
        """
        raw = payload.get("memo_ids")
        if not isinstance(raw, (list, tuple)) or not raw:
            return [], "memo_ids 得是一个非空的 id 数组"
        ids: list = []
        bad: list = []
        for item in raw:
            if isinstance(item, bool):        # bool 是 int 的子类，得先挡掉，否则 True 会变成 1
                bad.append(item)
                continue
            try:
                ids.append(int(item))
            except (TypeError, ValueError):
                bad.append(item)
        if bad:
            return [], f"这些不是合法的 id：{bad}"
        seen: set = set()
        unique: list = []
        for value in ids:
            if value not in seen:
                seen.add(value)
                unique.append(value)
        return unique, ""

    @app.post("/api/memos/merge")
    def api_memos_merge() -> Any:
        """🔗 合并选中项：把勾选的几段记录交给 DeepSeek 融会贯通成一张**新的**病历卡片。

        和「🔄 智能润色」是两回事：润色是在**同一张**卡片上做临床推演，合并是**另起一条**
        记录，把分几次记下的碎片熔成一份完整病历（``core.MERGE_PROMPT``）。

        三件事刻意这么定：

        * **源记录不删。** 合并不是"整理完就可以扔"的操作 —— 学生可能还要回头对照
          自己当时记的原话。真要清理，卡片上本来就有 🗑️ 删除，由他决定。
        * **先查 id 再谈合并。** 有一条不在（别的标签页刚删掉）就整个 404 报出来，
          而不是"能查到几条算几条"：少熔一段的病历看着一样完整，学生根本发现不了。
        * **按时间先后拼接**（``core.memo_sort_key``）。时间线是主诉里"3 天前"
          这类表述的锚点，顺序错了，模型会把后发生的事写成诱因。
        """
        payload = request.get_json(silent=True) or {}
        ids, err = _parse_memo_ids(payload)
        if err:
            return jsonify({"ok": False, "error": err}), 400
        if len(ids) < MERGE_MIN_SOURCES:
            return jsonify({
                "ok": False,
                "error": f"至少要选 {MERGE_MIN_SOURCES} 条记录才能合并"
                         f"（现在 {len(ids)} 条）；只合并一条的话，"
                         "该点那张卡片上的「🔄 智能润色/推演」",
            }), 400
        if len(ids) > MERGE_MAX_SOURCES:
            return jsonify({
                "ok": False,
                "error": f"一次最多合并 {MERGE_MAX_SOURCES} 条（现在 {len(ids)} 条）",
            }), 400

        sources = []
        missing = []
        for memo_id in ids:
            memo = core.fetch_memo(conn, memo_id)
            if memo is None:
                missing.append(memo_id)
            else:
                sources.append(memo)
        if missing:
            return jsonify({
                "ok": False,
                "error": f"这些卡片已经不在了：{missing}（可能在别的标签页被删过）",
                "missing": missing,
            }), 404

        target = app.config.get("PIPELINE")
        client = getattr(target, "client", None)
        if client is None:
            return jsonify({
                "ok": False,
                "error": "当前没有可用的 DeepSeek 客户端（没配 API Key，或用了 --no-llm），"
                         "合并整理需要它",
            }), 503

        sources.sort(key=core.memo_sort_key)          # 时间先后 —— 提示词里的时间线靠它
        started = time.time()
        try:
            merged = core.merge_memos(client, cfg.llm_model, sources)
        except Exception as exc:                       # 网络/解析失败都走这里
            return jsonify({"ok": False, "error": f"{type(exc).__name__}: {exc}"}), 502

        # 模型返回了合法 JSON，但八个病历字段全空 —— 那就是什么都没熔出来。
        # 这种"成功"的假象比报错更糟：页面上会多出一张空卡片，而学生以为合并好了。
        # 注意这里的判据是 core.has_soap（派生自 SOAP_SECTIONS），不是写死的五段 ——
        # 病历字段从五段扩到八段时，这条守卫自动跟着变。
        if not core.has_soap(merged):
            return jsonify({
                "ok": False,
                "error": "模型这次没返回任何病历内容（主诉到诊疗计划全空），已放弃本次合并、未入库",
            }), 502

        try:
            memo_id = core.insert_memo(conn, merged)
        except Exception as exc:                       # 不让界面 500
            return jsonify({"ok": False, "error": f"{type(exc).__name__}: {exc}"}), 500
        merged["id"] = memo_id

        synced = _refresh_markdown()
        # 新记录走 publish_memo：它会 +1 计数并广播，页面上"已记录 N 条"随之变化。
        state.publish_memo(merged)
        state.publish_announce(merged, "merge")   # 合并完也念一遍新病历的结论
        return jsonify({
            "ok": True,
            "id": memo_id,
            "memo": core.memo_to_dict(merged),
            "merged": [m.get("id") for m in sources],
            "record_time": merged.get("record_time") or "",
            "elapsed_ms": int((time.time() - started) * 1000),
            "markdown_path": cfg.markdown_path or "",
            "markdown_synced": synced,
        })

    def _refresh_markdown() -> bool:
        """按数据库现状整份重刷 ``my_memos.md``，返回是否真的同步成功。

        写失败不该让接口失败，但**也不能假装成功** —— 页面上要如实显示
        "数据库已存、文档没跟上"，否则学生导出的文档里少一条却毫无察觉。
        """
        if not cfg.markdown_path:
            return False
        try:
            core.export_markdown(conn, cfg.markdown_path)
            return True
        except Exception as exc:                        # 文档写失败不该让接口失败
            core.log(f"重刷 Markdown 失败：{type(exc).__name__}: {exc}", "导出")
            return False

    def _sync_markdown_and_broadcast(memo: dict) -> bool:
        """改完一条记录后的统一收尾：重刷 Markdown + 把新内容推给所有打开的页面。

        走 ``publish_update`` 而不是 ``publish_memo`` —— 改一条不是存一条，
        计数一涨页面上那句"已记录 N 条"就骗人了。**新**卡片请走 ``publish_memo``。
        """
        synced = _refresh_markdown()
        state.publish_update(memo)
        return synced

    @app.post("/api/control/<action>")
    def api_control(action: str) -> Any:
        """网页上的「结束并整理 / 放弃本条」两个按钮。

        只置一个标志位就返回，真正的收尾（截断音频、转写、落库或丢弃）由录音主循环
        下一轮完成 —— Web 线程绝不直接碰音频缓冲，否则会跟 VAD 状态机打架。
        """
        if action not in ("finish", "cancel"):
            return jsonify({"ok": False, "error": "action 只能是 finish 或 cancel"}), 400
        target = app.config.get("LISTENER")
        if target is None:
            return jsonify({"ok": False,
                            "error": "当前进程没有开麦克风（--no-mic），没有可控制的录音"}), 503
        accepted = (target.request_finish() if action == "finish"
                    else target.request_cancel())
        if not accepted:
            return jsonify({
                "ok": False,
                "state": target.state,
                "error": "现在没有在录音（可能已经整理完，或还在待机）",
            }), 409
        return jsonify({"ok": True, "action": action, "state": target.state})

    @app.post("/api/stop")
    def api_stop() -> Any:
        """停止监听（等同语音说"退出录制"）。"""
        target = app.config.get("LISTENER")
        if target is None:
            return jsonify({"ok": False, "error": "当前没有在监听"}), 503
        target.exit_requested.set()
        target.stop()
        return jsonify({"ok": True})

    # ------------------------------------------------------------------ #
    # 导出下载
    # ------------------------------------------------------------------ #
    @app.get("/api/export")
    def api_export() -> Any:
        fmt = (request.args.get("fmt") or "md").lower()
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        if fmt in ("md", "markdown"):
            memos = core.fetch_all_memos(conn)
            subtitle = (f"共 {len(memos)} 条记录 · 导出时间 "
                        f"{datetime.now():%Y-%m-%d %H:%M:%S}")
            body = "".join(
                core.format_markdown_entry(m, index=i) for i, m in enumerate(memos, 1)
            )
            data = (core._markdown_header(core.MARKDOWN_DOC_TITLE, subtitle) + body)
            return send_file(
                io.BytesIO(data.encode("utf-8")),
                mimetype="text/markdown",       # Flask 会自动补 charset=utf-8
                as_attachment=True,
                download_name=f"memos-{stamp}.md",
            )
        if fmt in ("docx", "word"):
            try:
                from docx import Document
            except ImportError:
                return jsonify({
                    "ok": False,
                    "error": "导出 Word 需要 python-docx，请先安装：pip install python-docx",
                }), 501
            memos = core.fetch_all_memos(conn)
            document = Document()
            document.add_heading(core.MARKDOWN_DOC_TITLE, level=0)
            document.add_paragraph(
                f"由 voice-memo 自动生成 · 共 {len(memos)} 条记录 · "
                f"导出时间 {datetime.now():%Y-%m-%d %H:%M:%S}"
            )
            for index, memo in enumerate(memos, 1):
                document.add_heading(
                    f"{index}. {memo.get('title') or '未命名备忘录'}", level=2
                )
                document.add_paragraph(f"时间：{memo.get('record_time', '')}")
                document.add_paragraph(f"分类：{memo.get('category') or '其他'}")
                document.add_paragraph(f"摘要：{memo.get('summary') or '（无）'}")
                for todo in memo.get("todos") or []:
                    due = f"（截止 {todo['due']}）" if todo.get("due") else ""
                    document.add_paragraph(f"☐ {todo['task']}{due}", style="List Bullet")
                raw = " ".join((memo.get("raw_text") or "").split())
                if raw:
                    document.add_paragraph(f"原始识别原文：{raw}")
            buffer = io.BytesIO()
            document.save(buffer)
            buffer.seek(0)
            return send_file(
                buffer,
                mimetype="application/vnd.openxmlformats-officedocument"
                         ".wordprocessingml.document",
                as_attachment=True,
                download_name=f"memos-{stamp}.docx",
            )
        return jsonify({"ok": False, "error": "fmt 只能是 md 或 docx"}), 400

    return app


def set_listener(app: Any, listener: Any) -> None:
    """监听器晚于 Flask 创建时，用它补上（``/api/stop`` 需要）。"""
    app.config["LISTENER"] = listener


def _sse(event: dict) -> str:
    """把一条事件序列化成 SSE 帧。"""
    import json

    return f"data: {json.dumps(event, ensure_ascii=False)}\n\n"


def build_export_markdown(conn: sqlite3.Connection, core: Any) -> str:
    """整份 Markdown 文本（自检脚本直接断言内容，不用起 HTTP 服务）。"""
    memos = core.fetch_all_memos(conn)
    stamp = f"{datetime.now():%Y-%m-%d %H:%M:%S}"
    subtitle = f"共 {len(memos)} 条记录 · 导出时间 {stamp}"
    body = "".join(core.format_markdown_entry(m, index=i) for i, m in enumerate(memos, 1))
    return core._markdown_header(core.MARKDOWN_DOC_TITLE, subtitle) + body
